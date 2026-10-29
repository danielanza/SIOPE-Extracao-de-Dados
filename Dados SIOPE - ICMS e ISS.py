"""
Dados SIOPE - ICMS e ISS
=========================
Lê os CSVs de cache extraídos pelo script principal e gera a planilha
"Dados SIOPE - ICMS e ISS.xlsx" com duas abas:
  - Municipios : Cota-Parte do ICMS, ISS (todas as sub-métricas) e Compensacoes
  - Estados    : ICMS (todas as sub-métricas) e Compensacoes

Implementa a "Regra de extração.docx" versão 4 (28/10/2026):
  - Valor de cada conta: |vlr_receita_realizada| - |vlr_outras_deducoes|
  - ICMS dos Estados e Compensacoes: se a dedução passar de 10% da receita
    realizada da mesma conta, a dedução é ignorada (LIMIAR_DEDUCAO)
  - 2019-2020: deduções em conta própria (49111305000 no ISS; 49111302000 e
    49111302001 no ICMS, estas sujeitas ao limiar)
  - 2021+: multas e juros nas formas agrupada (dígitos 2 e 4) e separada (5 a 8)
  - Compensacoes: compensação das perdas de ICMS da LC 194/2022 (2023+), com
    regra de duplicidade entre a rubrica dos Estados e a dos Municípios

Fonte dos dados: pasta cache/ gerada pelo script "Extrair Todas Receitas - Todos os Entes.py"
Anos cobertos: 2019–2025; 2026 aplica as mesmas regras de 2023+
"""

from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# CAMINHOS
# ─────────────────────────────────────────────
BASE_DIR   = Path(__file__).parent
CACHE_DIR  = BASE_DIR / "Registro Receita - Todos Entes" / "cache"
LISTA_ENTES = BASE_DIR / "Lista de Entes.xlsx"
OUTPUT     = BASE_DIR / "Dados SIOPE - ICMS e ISS.xlsx"

# Dedução maior que este percentual da receita realizada é ignorada
# (vale para o ICMS dos Estados e para as Compensacoes — seções 4.1 e 5.4 da regra)
LIMIAR_DEDUCAO = 0.10

# ─────────────────────────────────────────────
# MAPEAMENTO DE CONTAS POR ANO
# Estrutura: { (ano_inicio, ano_fim): { metrica: [cod_conta, ...] } }
# Contas da mesma métrica são somadas.
# ─────────────────────────────────────────────

# Municípios — ISS
ISS = {
    (2019, 2020): {
        "ISS Principal":           ["41113050100"],
        "ISS Multas e Juros Mora": ["41911400000"],
        "ISS DA - Principal":      ["41931130000"],
        "ISS DA - Multas e Juros": ["41913130000"],
        "ISS FECOP":               ["41113050200"],
    },
    (2021, 2022): {
        "ISS Principal":           ["11180231"],
        "ISS Multas e Juros Mora": ["11180232", "11180235", "11180236"],   # agrupada + Multas + Juros
        "ISS DA - Principal":      ["11180233"],
        "ISS DA - Multas e Juros": ["11180234", "11180237", "11180238"],   # agrupada + Multas DA + Juros DA
        "ISS FECOP":               ["11180240"],
    },
    (2023, 2026): {
        "ISS Principal":           ["11145111"],
        "ISS Multas e Juros Mora": ["11145112", "11145115", "11145116"],
        "ISS DA - Principal":      ["11145113"],
        "ISS DA - Multas e Juros": ["11145114", "11145117", "11145118"],
        "ISS FECOP":               ["11145120"],
    },
}

# Municípios — ISS específico para o DF (2019-2020)
# O DF usa a conta agregada 41113050000 no lugar de 41113050100,
# e não possui ISS FECOP nesse período (ISS = soma métricas 2 a 5).
# Em 2021+ o DF usa as mesmas contas dos demais municípios.
ISS_DF = {
    (2019, 2020): {
        "ISS Principal":           ["41113050000"],   # Exclusivo DF (demais municípios: 41113050100)
        "ISS Multas e Juros Mora": ["41911400000"],
        "ISS DA - Principal":      ["41931130000"],
        "ISS DA - Multas e Juros": ["41913130000"],
        # sem ISS FECOP para o DF em 2019-2020
    },
    (2021, 2022): ISS[(2021, 2022)],
    (2023, 2026): ISS[(2023, 2026)],
}

# Municípios — Cota-Parte do ICMS (conta agregadora; não somar as filhas)
COTAPARTE = {
    (2019, 2020): {"Cota-Parte": ["41722010100"]},
    (2021, 2022): {"Cota-Parte": ["17280110"]},
    (2023, 2026): {"Cota-Parte": ["17215000"]},
}

# Estados — ICMS
ICMS = {
    (2019, 2020): {
        "ICMS Principal":           ["41113020100"],
        "ICMS Multas e Juros Mora": ["41911420000"],
        "ICMS DA - Principal":      ["41931150000"],
        "ICMS DA - Multas e Juros": ["41913150000"],
        "ICMS FECOP":               ["41113020200"],
    },
    (2021, 2022): {
        "ICMS Principal":           ["11180211"],
        "ICMS Multas e Juros Mora": ["11180212", "11180215", "11180216"],   # 11180215 = "Multas" (antes na DA)
        "ICMS DA - Principal":      ["11180213"],
        "ICMS DA - Multas e Juros": ["11180214", "11180217", "11180218"],   # 11180214 sem ocorrência na base
        "ICMS FECOP":               ["11180220"],
    },
    (2023, 2026): {
        "ICMS Principal":           ["11145011"],
        "ICMS Multas e Juros Mora": ["11145012", "11145015", "11145016"],
        "ICMS DA - Principal":      ["11145013"],
        "ICMS DA - Multas e Juros": ["11145014", "11145017", "11145018"],
        "ICMS FECOP":               ["11145020"],
    },
}

# Compensação das perdas de ICMS — LC 194/2022 (somente 2023+)
RUBRICA_EST = ["17196200", "17196300"]   # rubrica dos Estados (União → Estado)
RUBRICA_MUN = ["17295300"]               # rubrica dos Municípios (cota-parte de 25%)
COMPENSACOES = {
    (2023, 2026): {"Compensacoes": RUBRICA_MUN + RUBRICA_EST},
}

# Deduções em conta própria (2019-2020)
# Estrutura: { (ano_inicio, ano_fim): { metrica: (conta_deducao, conta_receita) } }
DED_ISS = {
    (2019, 2020): {"ISS Principal": ("49111305000", "41113050100")},
}
DED_ISS_DF = {
    (2019, 2020): {"ISS Principal": ("49111305000", "41113050000")},
}
DED_ICMS = {
    (2019, 2020): {
        "ICMS Principal": ("49111302000", "41113020100"),
        "ICMS FECOP":     ("49111302001", "41113020200"),
    },
}

# Ordem das colunas métricas nas abas (Compensacoes sempre por último)
COLUNAS_MUN = ["Cota-Parte", "ISS", "ISS Principal", "ISS Multas e Juros Mora",
               "ISS DA - Principal", "ISS DA - Multas e Juros", "ISS FECOP", "Compensacoes"]
COLUNAS_EST = ["ICMS", "ICMS Principal", "ICMS Multas e Juros Mora",
               "ICMS DA - Principal", "ICMS DA - Multas e Juros", "ICMS FECOP", "Compensacoes"]

ISS_SUB  = ["ISS Principal", "ISS Multas e Juros Mora", "ISS DA - Principal",
            "ISS DA - Multas e Juros", "ISS FECOP"]
ICMS_SUB = ["ICMS Principal", "ICMS Multas e Juros Mora", "ICMS DA - Principal",
            "ICMS DA - Multas e Juros", "ICMS FECOP"]

# ─────────────────────────────────────────────
# FUNÇÕES AUXILIARES
# ─────────────────────────────────────────────

def build_lookup(mapping_dict):
    """
    Retorna dict: { (cod_conta, ano): metrica }
    para lookup rápido ao processar cada linha.
    """
    lookup = {}
    for (ano_ini, ano_fim), contas in mapping_dict.items():
        for metrica, cods in contas.items():
            for cod in cods:
                for ano in range(ano_ini, ano_fim + 1):
                    lookup[(cod, ano)] = metrica
    return lookup

def valor_conta(realizada, deducao, limiar=None):
    """
    |realizada| - |deducao|. Com limiar: se a dedução passar de `limiar` da
    receita realizada, ela é ignorada e vale apenas |realizada|.
    """
    r, d = abs(realizada), abs(deducao)
    if limiar is not None and d > limiar * r:
        return r
    return r - d

def ler_cache(csv_path):
    dtype = {"sgl_uf": str, "cod_municipio": str, "num_ano": str, "cod_conta": str}
    df = pd.read_csv(csv_path, dtype=dtype)
    df = df.drop_duplicates()
    df["num_ano"] = pd.to_numeric(df["num_ano"], errors="coerce").astype("Int64")
    df["vlr_receita_realizada"] = pd.to_numeric(df["vlr_receita_realizada"], errors="coerce").fillna(0)
    df["vlr_outras_deducoes"]   = pd.to_numeric(df["vlr_outras_deducoes"],   errors="coerce").fillna(0)
    return df

def aplicar_duplicidade_compensacao(df, chave_grupo):
    """
    Regra de duplicidade (seção 5.2 da regra): no mesmo ente-ano, se a soma da
    rubrica dos Estados (17196200 + 17196300) for igual à rubrica dos Municípios
    (17295300), é o mesmo valor lançado duas vezes — mantém só a rubrica dos
    Municípios. Valores diferentes são tratados como parcelas distintas (somados).
    Retorna (df_filtrado, n_casos_duplicados).
    """
    comp = df[df["cod_conta"].isin(RUBRICA_EST + RUBRICA_MUN) & (df["num_ano"] >= 2023)]
    if comp.empty:
        return df, 0
    tot = comp.assign(rub=comp["cod_conta"].isin(RUBRICA_EST).map({True: "EST", False: "MUN"}),
                      v=comp["vlr_receita_realizada"].abs())
    tot = tot.groupby(chave_grupo + ["rub"])["v"].sum().unstack("rub").fillna(0)
    if "EST" not in tot or "MUN" not in tot:
        return df, 0
    dup = tot[(tot["EST"] > 0) & (tot["MUN"] > 0) & ((tot["EST"] - tot["MUN"]).abs() < 0.01)]
    if dup.empty:
        return df, 0
    chaves_dup = set(dup.index)
    idx_chave = pd.MultiIndex.from_frame(df[chave_grupo]) if len(chave_grupo) > 1 else df[chave_grupo[0]]
    remover = df["cod_conta"].isin(RUBRICA_EST) & pd.Series(idx_chave.isin(chaves_dup), index=df.index)
    return df[~remover], len(dup)

def calcular(df, chave_grupo, mapeamentos, deducoes=None, limiar_por_metrica=None):
    """
    Aplica os mapeamentos e devolve DataFrame pivotado (chave_grupo + métricas).
      mapeamentos        : lista de dicts {(ano_ini, ano_fim): {metrica: [contas]}}
      deducoes           : dict {(ano_ini, ano_fim): {metrica: (conta_ded, conta_receita)}}
      limiar_por_metrica : {metrica: limiar} — métricas sem entrada não têm limiar
    Também devolve estatísticas de deduções ignoradas pelo limiar.
    """
    limiar_por_metrica = limiar_por_metrica or {}
    lookup = {}
    for mp in mapeamentos:
        lookup.update(build_lookup(mp))
    contas = {cod for (cod, _) in lookup}

    base = df[df["cod_conta"].isin(contas)].copy()
    base["_metrica"] = [lookup.get((c, a)) for c, a in zip(base["cod_conta"], base["num_ano"])]
    base = base[base["_metrica"].notna()].copy()
    base["_valor"] = [
        valor_conta(r, d, limiar_por_metrica.get(m))
        for r, d, m in zip(base["vlr_receita_realizada"], base["vlr_outras_deducoes"], base["_metrica"])
    ]
    ignoradas = int(sum(
        1 for r, d, m in zip(base["vlr_receita_realizada"], base["vlr_outras_deducoes"], base["_metrica"])
        if limiar_por_metrica.get(m) is not None and abs(d) > 0 and abs(d) > limiar_por_metrica[m] * abs(r)
    ))
    partes = [base[chave_grupo + ["_metrica", "_valor"]]]

    # Deduções em conta própria (2019-2020)
    ded_aplicadas = ded_ignoradas = 0
    for (ano_ini, ano_fim), regras in (deducoes or {}).items():
        anos = df["num_ano"].between(ano_ini, ano_fim).fillna(False).astype(bool)
        for metrica, (cod_ded, cod_rec) in regras.items():
            d = df[anos & (df["cod_conta"] == cod_ded)].groupby(chave_grupo)["vlr_receita_realizada"].apply(lambda s: s.abs().sum())
            if d.empty:
                continue
            r = df[anos & (df["cod_conta"] == cod_rec)].groupby(chave_grupo)["vlr_receita_realizada"].apply(lambda s: s.abs().sum())
            r = r.reindex(d.index).fillna(0)
            lim = limiar_por_metrica.get(metrica)
            ok = (d > 0) if lim is None else ((d > 0) & ~(d > lim * r))
            ded_aplicadas += int(ok.sum())
            ded_ignoradas += int(((d > 0) & ~ok).sum())
            aplicar = (-d[ok]).rename("_valor").reset_index()
            aplicar["_metrica"] = metrica
            partes.append(aplicar[chave_grupo + ["_metrica", "_valor"]])

    longo = pd.concat(partes, ignore_index=True)
    if longo.empty:
        return pd.DataFrame(), {}
    pivot = longo.pivot_table(index=chave_grupo, columns="_metrica", values="_valor", aggfunc="sum").reset_index()
    pivot.columns.name = None
    stats = {"linhas_ignoradas_limiar": ignoradas, "deducoes_conta_aplicadas": ded_aplicadas,
             "deducoes_conta_ignoradas": ded_ignoradas}
    return pivot, stats

# ─────────────────────────────────────────────
# CARREGAR LISTA DE ENTES (cod_ibge dos estados)
# ─────────────────────────────────────────────

def carregar_cod_ibge_estados():
    wb = openpyxl.load_workbook(LISTA_ENTES, read_only=True)
    ws = wb["Estados"]
    mapa = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] and row[2]:
            mapa[str(row[2]).strip()] = str(int(row[1]))
    wb.close()
    return mapa

# ─────────────────────────────────────────────
# PROCESSAR MUNICÍPIOS
# ─────────────────────────────────────────────

LIMIAR_MUN = {"Compensacoes": LIMIAR_DEDUCAO}

def processar_df_como_municipio():
    """
    O DF declara como Estado, então seus dados estão em est_todos.csv.
    Extrai as linhas do DF e aplica ISS (contas do DF) e Cota-Parte.
    A compensação do DF é recebida como Estado e entra na aba Estados.
    """
    csv_path = CACHE_DIR / "est_todos.csv"
    if not csv_path.exists():
        return pd.DataFrame()
    df = ler_cache(csv_path)
    df = df[df["sgl_uf"] == "DF"].copy()
    if df.empty:
        return pd.DataFrame()
    df["cod_municipio"] = "530010"
    pivot, _ = calcular(df, ["sgl_uf", "cod_municipio", "num_ano"], [ISS_DF, COTAPARTE], deducoes=DED_ISS_DF)
    return pivot

def processar_municipios():
    csvs = sorted(CACHE_DIR.glob("mun_*.csv"))
    partes = []
    total = len(csvs)
    chave = ["sgl_uf", "cod_municipio", "num_ano"]
    resumo = {"dup": 0, "ign": 0, "ded": 0}

    for i, csv_path in enumerate(csvs, 1):
        uf = csv_path.stem.replace("mun_", "")
        print(f"  [{i:02d}/{total}] Municipios {uf}...", end=" ")

        df = ler_cache(csv_path)
        df, n_dup = aplicar_duplicidade_compensacao(df, chave)
        pivot, st = calcular(df, chave, [ISS, COTAPARTE, COMPENSACOES], deducoes=DED_ISS, limiar_por_metrica=LIMIAR_MUN)

        if pivot.empty:
            print("sem dados relevantes")
            continue

        resumo["dup"] += n_dup
        resumo["ign"] += st["linhas_ignoradas_limiar"]
        resumo["ded"] += st["deducoes_conta_aplicadas"]
        print(f"{len(pivot):,} linhas")
        partes.append(pivot)

    # DF: dados municipais estão no CSV de estados (DF é Município e Estado)
    print(f"  [DF] Municipio DF (via est_todos.csv)...", end=" ")
    pivot_df = processar_df_como_municipio()
    if not pivot_df.empty:
        print(f"{len(pivot_df):,} linhas")
        partes.append(pivot_df)
    else:
        print("sem dados relevantes")

    print(f"  Compensacoes: {resumo['dup']} duplicidade(s) entre rubricas consolidada(s); "
          f"{resumo['ign']} dedução(ões) ignorada(s) pelo limiar")
    print(f"  ISS 2019-2020: dedução 49111305000 aplicada em {resumo['ded']:,} município-ano")

    if not partes:
        return pd.DataFrame()

    df = pd.concat(partes, ignore_index=True)

    # ISS = soma das métricas 2 a 6 (Compensacoes não entra no ISS)
    for col in ISS_SUB:
        if col not in df.columns:
            df[col] = 0.0
    df["ISS"] = df[ISS_SUB].fillna(0).sum(axis=1)

    # Garantir todas as colunas métricas
    for col in COLUNAS_MUN:
        if col not in df.columns:
            df[col] = None

    df = df.rename(columns={"sgl_uf": "UF", "cod_municipio": "COD_IBGE", "num_ano": "ANO"})
    df.insert(3, "BASE", "SIOPE")
    df = df[["UF", "COD_IBGE", "ANO", "BASE"] + COLUNAS_MUN]
    df = df.sort_values(["UF", "COD_IBGE", "ANO"]).reset_index(drop=True)
    return df

# ─────────────────────────────────────────────
# PROCESSAR ESTADOS
# ─────────────────────────────────────────────

LIMIAR_EST = {m: LIMIAR_DEDUCAO for m in ICMS_SUB + ["Compensacoes"]}

def processar_estados():
    csv_path = CACHE_DIR / "est_todos.csv"
    print(f"  Estados...", end=" ")

    chave = ["sgl_uf", "num_ano"]
    df = ler_cache(csv_path)
    df, n_dup = aplicar_duplicidade_compensacao(df, chave)
    pivot, st = calcular(df, chave, [ICMS, COMPENSACOES], deducoes=DED_ICMS, limiar_por_metrica=LIMIAR_EST)

    if pivot.empty:
        print("sem dados relevantes")
        return pd.DataFrame()

    print(f"{len(pivot):,} linhas")
    print(f"  Limiar de {LIMIAR_DEDUCAO:.0%}: {st['linhas_ignoradas_limiar']} dedução(ões) em coluna ignorada(s); "
          f"2019-2020 em conta própria: {st['deducoes_conta_aplicadas']} aplicada(s), {st['deducoes_conta_ignoradas']} ignorada(s)")
    print(f"  Compensacoes: {n_dup} duplicidade(s) entre rubricas consolidada(s)")

    # ICMS = soma das métricas 2 a 6 (Compensacoes não entra no ICMS)
    for col in ICMS_SUB:
        if col not in pivot.columns:
            pivot[col] = 0.0
    pivot["ICMS"] = pivot[ICMS_SUB].fillna(0).sum(axis=1)

    for col in COLUNAS_EST:
        if col not in pivot.columns:
            pivot[col] = None

    cod_ibge_map = carregar_cod_ibge_estados()
    pivot = pivot.rename(columns={"sgl_uf": "UF", "num_ano": "ANO"})
    pivot.insert(1, "COD_IBGE", pivot["UF"].map(cod_ibge_map))
    pivot.insert(3, "BASE", "SIOPE")
    pivot = pivot[["UF", "COD_IBGE", "ANO", "BASE"] + COLUNAS_EST]
    pivot = pivot.sort_values(["UF", "ANO"]).reset_index(drop=True)
    return pivot

# ─────────────────────────────────────────────
# ESCREVER EXCEL
# ─────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")

def _escrever_aba(wb, nome_aba, df):
    ws = wb.create_sheet(nome_aba)

    # Cabeçalho
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    # Dados (NaN vira célula vazia)
    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        for col_idx, val in enumerate(row, 1):
            if val is not None and not (isinstance(val, float) and pd.isna(val)) and val is not pd.NA:
                ws.cell(row=row_idx, column=col_idx, value=val)

    # Larguras
    larguras = {"UF": 6, "COD_IBGE": 12, "ANO": 7, "BASE": 8}
    for col_idx, col_name in enumerate(df.columns, 1):
        w = larguras.get(col_name, 24)
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

# ─────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────

def main():
    print("=== Dados SIOPE - ICMS e ISS ===\n")

    print("Processando municipios...")
    df_mun = processar_municipios()
    print(f"  Total municipios: {len(df_mun):,} linhas\n")

    print("Processando estados...")
    df_est = processar_estados()
    print(f"  Total estados: {len(df_est):,} linhas\n")

    print(f"Gerando {OUTPUT.name}...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove aba padrão

    if not df_mun.empty:
        _escrever_aba(wb, "Municipios", df_mun)
    if not df_est.empty:
        _escrever_aba(wb, "Estados", df_est)

    wb.save(OUTPUT)
    print(f"Arquivo salvo: {OUTPUT}")

if __name__ == "__main__":
    main()
